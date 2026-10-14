import openpyxl
import json

wb = openpyxl.load_workbook(r'C:\Users\11828\Desktop\CET4_简化清晰版_每日单词已加入_最终版.xlsx', data_only=True)

# 1. 从「每日计划」提取完整数据
ws_plan = wb['每日计划']
plan_data = {}
status_set = set()  # 收集所有出现过的状态符号

for row in range(2, ws_plan.max_row + 1):
    day_num = ws_plan.cell(row=row, column=1).value
    date_val = ws_plan.cell(row=row, column=2).value
    if not day_num or not date_val:
        continue
    
    if hasattr(date_val, 'strftime'):
        date_str = date_val.strftime('%Y-%m-%d')
    else:
        date_str = str(date_val)[:10]
    
    task = ws_plan.cell(row=row, column=6).value or ''
    word_count = ws_plan.cell(row=row, column=7).value or 12
    review = ws_plan.cell(row=row, column=8).value or ''
    duration = ws_plan.cell(row=row, column=9).value or ''
    status = ws_plan.cell(row=row, column=10).value or ''
    
    status_set.add(str(status).strip())
    
    plan_data[int(day_num)] = {
        'date': date_str,
        'task': str(task),
        'review': str(review),
        'duration': str(duration),
        'taskStatus': str(status).strip(),
        'word_count': int(word_count)
    }

print('Excel 中出现过的完成状态符号：', status_set)
print()

# 2. 从「词汇总表」提取每天的单词
ws_words = wb['词汇总表']
words_by_day = {}
for row in range(2, ws_words.max_row + 1):
    day_num = ws_words.cell(row=row, column=1).value
    en = ws_words.cell(row=row, column=3).value
    cn = ws_words.cell(row=row, column=5).value
    if not day_num or not en:
        continue
    day_num = int(day_num)
    if day_num not in words_by_day:
        words_by_day[day_num] = []
    words_by_day[day_num].append({'en': str(en), 'cn': str(cn) if cn else ''})

# 3. 合并数据，生成 JS 数组
result = []
for day_num in sorted(plan_data.keys()):
    plan = plan_data[day_num]
    words = words_by_day.get(day_num, [])
    result.append({
        'date': plan['date'],
        'dayNum': day_num,
        'wordCount': len(words),
        'task': plan['task'],
        'review': plan['review'],
        'duration': plan['duration'],
        'taskStatus': plan['taskStatus'],
        'words': words
    })

# 4. 生成 JS 文件内容
js_content = 'var DAILY_PLAN_DATA = ' + json.dumps(result, ensure_ascii=False, indent=2) + ';\n'

# 5. 写入文件
with open(r'D:\Users\11828\Documents\GitHub\cet4-dictation\js\daily_plan.js', 'w', encoding='utf-8') as f:
    f.write(js_content)

print(f'成功生成 daily_plan.js，共 {len(result)} 天')
print(f'第 1 天状态: {result[0]["taskStatus"]}')
print(f'第 2 天状态: {result[1]["taskStatus"]}')
print(f'第 1 天任务内容长度: {len(result[0]["task"])} 字')
