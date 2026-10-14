import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\11828\Desktop\CET4_简化清晰版_每日单词已加入_最终版.xlsx', data_only=True)

# 检查「每日计划」工作表的完整列结构
ws = wb['每日计划']
print('=== 每日计划 工作表结构 ===')
print(f'共 {ws.max_row} 行, {ws.max_column} 列')
print()

# 打印表头（第1行）
print('表头：')
for col in range(1, ws.max_column + 1):
    cell = ws.cell(row=1, column=col)
    print(f'  列{col}: {cell.value}')

print()
# 打印第2行（第1天）的所有列
print('第1天（行2）的所有列：')
for col in range(1, ws.max_column + 1):
    cell = ws.cell(row=2, column=col)
    val = str(cell.value)[:50] if cell.value else '(空)'
    print(f'  列{col}: {val}')

print()
# 看看有没有状态列，打印第2-5行的最后几列
print('第2-5行的最后5列：')
for row in range(2, 6):
    row_data = []
    for col in range(max(1, ws.max_column - 4), ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        row_data.append(str(cell.value)[:20] if cell.value else '')
    print(f'  行{row}: {row_data}')
